import argparse
import json
import os
import re
import sys
import time
import urllib.error
import urllib.request
from datetime import date, datetime, timezone as datetime_timezone, timedelta
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from io import BytesIO
from pathlib import Path
from threading import Thread
from zoneinfo import ZoneInfo

import requests
from openpyxl import Workbook
from openpyxl import load_workbook


if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = Path(os.getenv("DATA_DIR", BASE_DIR / "data"))
CONFIG_PATH = BASE_DIR / "config.local.json"
CHAT_IDS_PATH = Path(os.getenv("CHAT_IDS_PATH", DATA_DIR / "chat_ids.json"))
PENDING_ATTENDANCE_PATH = Path(os.getenv("PENDING_ATTENDANCE_PATH", DATA_DIR / "pending_attendance.json"))
ABSENCES_FALLBACK_PATH = Path(os.getenv("ABSENCES_FALLBACK_PATH", DATA_DIR / "absences_fallback.json"))
ABSENCES_XLSX_PATH = Path(os.getenv("ABSENCES_XLSX_PATH", DATA_DIR / "absences.xlsx"))

# Main operational settings. Change these values, then restart the bot.
# Leave these empty to read year/month from the class sheet title.
SCHEDULE_YEAR_OVERRIDE = ""
SCHEDULE_MONTH_OVERRIDE = ""
CURRENT_JALALI_DATE_OVERRIDE = ""
REMINDER_TIME = "09:00"
ATTENDANCE_TIME = "23:07"
SCHEDULER_INTERVAL_SECONDS = 15
BOT_TIMEZONE = "Asia/Tehran"
MANUAL_TEST_CLASSES = []

DEFAULT_SHEET_EXPORT_URL = (
    # Active class schedule sheet:
    # https://docs.google.com/spreadsheets/d/1jwQ-2k6zbOGLTgPjSpOXGvnlBRWm70Mk/edit
    "https://docs.google.com/spreadsheets/d/"
    "1jwQ-2k6zbOGLTgPjSpOXGvnlBRWm70Mk/export?format=xlsx"
)
DEFAULT_CONTACTS_EXPORT_URL = (
    # Professor phone sheet + absence sheet:
    # https://docs.google.com/spreadsheets/d/1P_wWkcMIpsUZYME8xCllQRvjfHSGCa0s/edit
    "https://docs.google.com/spreadsheets/d/"
    "1P_wWkcMIpsUZYME8xCllQRvjfHSGCa0s/export?format=xlsx"
)
DEFAULT_ABSENCE_WEBHOOK_URL = (
    "https://script.google.com/macros/s/"
    "AKfycbw6Mn9mGMjkUsQKaLRiY6MDV22Xc6jtWB4BPpzJo3vQk7rvr1wC6h-ZfQQwD89FECo/exec"
)
START_MESSAGE = (
    "با عرض سلام و خوش آمد خدمت اساتید گرامی\n"
    "لطفا شماره همراه خود را ( با کیبورد انگلیسی ) جهت مشاهده ی نام دانشجویان اینترن مربوط به مطب خود را وارد نمایید"
)

INVALID_PHONE_MESSAGE = "با عرض پوزش شماره ی وارد شده ثبت نشده است"
SCHEDULE_FOOTER = "برنامه کلاس های شما به شکل بالا هست و در روز کلاس برای شما یک پیام یاداوری ارسال خواهد شد"
VIEW_CLASSES_TEXT = "برنامه ماهانه کلینیک ویژه من"
ATTENDANCE_YES_TEXT = "بله"
ATTENDANCE_NO_TEXT = "خیر"
REGISTERED_PHONE_MESSAGE = "شماره همراه شما قبلا با شماره {phone} ثبت شده است."
LOGIN_SUCCESS_MESSAGE = (
    "با عرض سلام خدمت {professor_name}\n"
    "ضمن تشکر از زحمات شما\n"
    "لیست دانشجویان ماه آینده ی شما به صورت زیر است"
)

PERSIAN_MONTHS = {
    "فروردین": 1,
    "اردیبهشت": 2,
    "خرداد": 3,
    "تیر": 4,
    "مرداد": 5,
    "شهریور": 6,
    "مهر": 7,
    "آبان": 8,
    "آذر": 9,
    "دی": 10,
    "بهمن": 11,
    "اسفند": 12,
}

PERSIAN_WEEKDAYS = [
    "دوشنبه",
    "سه شنبه",
    "چهارشنبه",
    "پنج شنبه",
    "جمعه",
    "شنبه",
    "یکشنبه",
]

REMINDERS_SENT_THIS_RUN = set()
ATTENDANCE_SENT_THIS_RUN = set()


def read_json(path, default):
    try:
        with path.open("r", encoding="utf-8") as file:
            return json.load(file)
    except FileNotFoundError:
        return default
    except json.JSONDecodeError as exc:
        raise RuntimeError(f"Invalid JSON in {path}: {exc}") from exc


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as file:
        json.dump(value, file, ensure_ascii=False, indent=2)
        file.write("\n")


def normalize_text(value):
    text = str(value or "").strip()
    text = text.replace("ي", "ی").replace("ك", "ک")
    return re.sub(r"\s+", " ", text)


def normalize_phone(value):
    if isinstance(value, float) and value.is_integer():
        value = str(int(value))
    elif isinstance(value, int):
        value = str(value)
    else:
        value = str(value or "").strip()
        if re.fullmatch(r"\d+\.0", value):
            value = value[:-2]
    digits = re.sub(r"\D+", "", value)
    if digits.startswith("0098"):
        digits = "0" + digits[4:]
    elif digits.startswith("98") and len(digits) == 12:
        digits = "0" + digits[2:]
    elif digits.startswith("9") and len(digits) == 10:
        digits = "0" + digits
    return digits


def compact_text(value):
    return normalize_text(value).replace(" ", "")


def to_persian_digits(value):
    return str(value).translate(str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹"))


def gregorian_to_jalali(g_year, g_month, g_day):
    g_days_in_month = [31, 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]
    j_days_in_month = [31, 31, 31, 31, 31, 31, 30, 30, 30, 30, 30, 29]
    gy = g_year - 1600
    gm = g_month - 1
    gd = g_day - 1
    g_day_no = 365 * gy + (gy + 3) // 4 - (gy + 99) // 100 + (gy + 399) // 400
    for index in range(gm):
        g_day_no += g_days_in_month[index]
    leap = (gy + 1600) % 4 == 0 and ((gy + 1600) % 100 != 0 or (gy + 1600) % 400 == 0)
    if gm > 1 and leap:
        g_day_no += 1
    g_day_no += gd
    j_day_no = g_day_no - 79
    j_np = j_day_no // 12053
    j_day_no %= 12053
    jy = 979 + 33 * j_np + 4 * (j_day_no // 1461)
    j_day_no %= 1461
    if j_day_no >= 366:
        jy += (j_day_no - 1) // 365
        j_day_no = (j_day_no - 1) % 365
    jm = 0
    while jm < 11 and j_day_no >= j_days_in_month[jm]:
        j_day_no -= j_days_in_month[jm]
        jm += 1
    return jy, jm + 1, j_day_no + 1


def jalali_to_gregorian(j_year, j_month, j_day):
    j_days_in_month = [31, 31, 31, 31, 31, 31, 30, 30, 30, 30, 30, 29]
    g_days_in_month = [31, 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]

    jy = j_year - 979
    jm = j_month - 1
    jd = j_day - 1

    j_day_no = 365 * jy + (jy // 33) * 8 + ((jy % 33) + 3) // 4
    for index in range(jm):
        j_day_no += j_days_in_month[index]
    j_day_no += jd

    g_day_no = j_day_no + 79
    gy = 1600 + 400 * (g_day_no // 146097)
    g_day_no %= 146097

    leap = True
    if g_day_no >= 36525:
        g_day_no -= 1
        gy += 100 * (g_day_no // 36524)
        g_day_no %= 36524
        if g_day_no >= 365:
            g_day_no += 1
        else:
            leap = False

    gy += 4 * (g_day_no // 1461)
    g_day_no %= 1461

    if g_day_no >= 366:
        leap = False
        g_day_no -= 1
        gy += g_day_no // 365
        g_day_no %= 365

    gm = 0
    while gm < 11:
        days_in_month = g_days_in_month[gm]
        if gm == 1 and leap:
            days_in_month += 1
        if g_day_no < days_in_month:
            break
        g_day_no -= days_in_month
        gm += 1

    return gy, gm + 1, g_day_no + 1


def jalali_weekday_name(j_year, j_month, j_day):
    g_year, g_month, g_day = jalali_to_gregorian(j_year, j_month, j_day)
    return PERSIAN_WEEKDAYS[date(g_year, g_month, g_day).weekday()]


def parse_jalali_date(value):
    if value:
        year, month, day = [int(part) for part in value.split("-")]
        return year, month, day
    today = date.today()
    return gregorian_to_jalali(today.year, today.month, today.day)


def config_value(config, env_name, key, default=None):
    value = os.getenv(env_name)
    if value is not None:
        return value
    return config.get(key, default)


def config_bool(config, env_name, key, default=False):
    value = config_value(config, env_name, key, default)
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in {"1", "true", "yes", "on"}


def parse_time(value, default):
    value = str(value or default).strip()
    if ":" in value:
        hour, minute = value.split(":", 1)
        return int(hour), int(minute)
    return int(value), 0


def time_reached(now, value):
    hour, minute = parse_time(value, "00:00")
    return (now.hour, now.minute) >= (hour, minute)


def time_matches(now, value):
    hour, minute = parse_time(value, "00:00")
    return now.hour == hour and now.minute == minute


def get_bot_timezone(timezone_name):
    timezone_name = str(timezone_name or BOT_TIMEZONE).strip() or BOT_TIMEZONE
    try:
        return ZoneInfo(timezone_name)
    except Exception as exc:
        if timezone_name == "Asia/Tehran":
            print(f"Timezone data for {timezone_name} is unavailable, using fixed +03:30 offset: {exc}")
            return datetime_timezone(timedelta(hours=3, minutes=30), timezone_name)
        print(f"Timezone data for {timezone_name} is unavailable, using UTC: {exc}")
        return datetime_timezone.utc


def timezone_label(timezone):
    return getattr(timezone, "key", str(timezone))


def load_config():
    config = read_json(CONFIG_PATH, {})
    token = os.getenv("BOT_TOKEN", config.get("bot_token", "")).strip()
    if not token:
        raise RuntimeError("Set BOT_TOKEN environment variable before running the bot.")
    config["bot_token"] = token
    config["api_base_url"] = os.getenv("API_BASE_URL", config.get("api_base_url", "https://tapi.bale.ai/bot"))
    config["poll_timeout_seconds"] = os.getenv("POLL_TIMEOUT_SECONDS", config.get("poll_timeout_seconds", 25))
    config["sheet_export_url"] = os.getenv("SHEET_EXPORT_URL", config.get("sheet_export_url", DEFAULT_SHEET_EXPORT_URL))
    config["contacts_export_url"] = os.getenv(
        "CONTACTS_EXPORT_URL",
        config.get("contacts_export_url", DEFAULT_CONTACTS_EXPORT_URL),
    )
    config["absence_webhook_url"] = os.getenv(
        "ABSENCE_WEBHOOK_URL",
        config.get("absence_webhook_url", DEFAULT_ABSENCE_WEBHOOK_URL),
    )
    config["current_jalali_date"] = os.getenv("CURRENT_JALALI_DATE", CURRENT_JALALI_DATE_OVERRIDE)
    config["schedule_year_override"] = os.getenv("SCHEDULE_YEAR_OVERRIDE", str(SCHEDULE_YEAR_OVERRIDE))
    config["schedule_month_override"] = os.getenv("SCHEDULE_MONTH_OVERRIDE", str(SCHEDULE_MONTH_OVERRIDE))
    config["reminder_time"] = os.getenv("REMINDER_TIME", REMINDER_TIME)
    config["attendance_time"] = os.getenv("ATTENDANCE_TIME", ATTENDANCE_TIME)
    config["scheduler_interval_seconds"] = int(os.getenv("SCHEDULER_INTERVAL_SECONDS", str(SCHEDULER_INTERVAL_SECONDS)))
    config["bot_timezone"] = os.getenv("BOT_TIMEZONE", BOT_TIMEZONE)
    return config


class HealthHandler(BaseHTTPRequestHandler):
    def do_GET(self):
        self.send_response(200)
        self.send_header("Content-Type", "text/plain; charset=utf-8")
        self.end_headers()
        self.wfile.write("ok\n".encode("utf-8"))

    def log_message(self, format, *args):
        return


def start_health_server():
    port = int(os.getenv("PORT", "8000"))
    server = ThreadingHTTPServer(("0.0.0.0", port), HealthHandler)
    Thread(target=server.serve_forever, daemon=True).start()
    print(f"Health server is listening on port {port}.")


class BaleBot:
    def __init__(self, config):
        token = config["bot_token"].strip()
        api_base_url = config.get("api_base_url", "https://tapi.bale.ai/bot").rstrip("/")
        self.base_url = f"{api_base_url}{token}"
        self.poll_timeout = int(config.get("poll_timeout_seconds", 25))
        self.offset = None

    def request(self, method, payload=None):
        url = f"{self.base_url}/{method}"
        data = None
        headers = {}
        if payload is not None:
            data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
            headers["Content-Type"] = "application/json"
        request = urllib.request.Request(url, data=data, headers=headers, method="POST")
        with urllib.request.urlopen(request, timeout=self.poll_timeout + 10) as response:
            return json.loads(response.read().decode("utf-8"))

    def get_updates(self):
        payload = {"timeout": self.poll_timeout, "limit": 50}
        if self.offset is not None:
            payload["offset"] = self.offset
        return self.request("getUpdates", payload)

    def send_message(self, chat_id, text, keyboard=None):
        payload = {"chat_id": chat_id, "text": text}
        if keyboard:
            payload["reply_markup"] = keyboard
        return self.request("sendMessage", payload)


class WorkbookCache:
    def __init__(self, export_url):
        self.export_url = export_url
        self.loaded_at = 0
        self.cache_seconds = int(os.getenv("SHEET_CACHE_SECONDS", "300"))
        self.workbook = None

    def workbook_data(self):
        now = time.time()
        if self.workbook is None or now - self.loaded_at > self.cache_seconds:
            response = requests.get(self.export_url, timeout=30)
            response.raise_for_status()
            self.workbook = load_workbook(BytesIO(response.content), data_only=True)
            self.loaded_at = now
        return self.workbook


class SheetSchedule:
    def __init__(self, export_url, config=None):
        self.config = config or {}
        self.cache = WorkbookCache(export_url)

    def worksheet(self):
        workbook = self.cache.workbook_data()
        return workbook[workbook.sheetnames[0]]

    def month_info(self):
        year_override = config_value(self.config, "SCHEDULE_YEAR_OVERRIDE", "schedule_year_override")
        month_override = config_value(self.config, "SCHEDULE_MONTH_OVERRIDE", "schedule_month_override")
        if year_override and month_override:
            month_number = int(month_override)
            month_name = next(name for name, number in PERSIAN_MONTHS.items() if number == month_number)
            return int(year_override), month_number, month_name

        ws = self.worksheet()
        title = " ".join(normalize_text(cell.value) for cell in ws[1] if cell.value)
        found_month = None
        for month_name, month_number in PERSIAN_MONTHS.items():
            if month_name in title:
                found_month = (month_name, month_number)
                break
        year_match = re.search(r"(\d{4})", title)
        year = int(year_match.group(1)) if year_match else None
        if found_month and year:
            return year, found_month[1], found_month[0]
        current_year, current_month, _ = parse_jalali_date(self.config.get("current_jalali_date"))
        month_name = next(name for name, number in PERSIAN_MONTHS.items() if number == current_month)
        return year or current_year, current_month, month_name

    def professor_columns(self):
        ws = self.worksheet()
        columns = {}
        for col in range(1, ws.max_column + 1):
            value = normalize_text(ws.cell(2, col).value)
            if value.startswith("دکتر"):
                columns[value] = col
        return columns

    def schedule_for_professor(self, professor_name):
        ws = self.worksheet()
        professor_name = normalize_text(professor_name)
        schedule_year, schedule_month, _ = self.month_info()
        columns = self.professor_columns()
        col = columns.get(professor_name)
        if not col:
            compact_columns = {compact_text(name): col for name, col in columns.items()}
            col = compact_columns.get(compact_text(professor_name))
        if not col:
            return []
        classes = []
        for row in range(3, ws.max_row + 1):
            weekday = normalize_text(ws.cell(row, 1).value)
            day_number = ws.cell(row, 2).value
            student = normalize_text(ws.cell(row, col).value)
            if not weekday or not day_number or not student or student == "*":
                continue
            day_number = int(day_number)
            classes.append(
                {
                    "weekday": jalali_weekday_name(schedule_year, schedule_month, day_number),
                    "day": day_number,
                    "student": student,
                }
            )
        professor_key = compact_text(professor_name)
        for item in MANUAL_TEST_CLASSES:
            if compact_text(item.get("professor")) != professor_key:
                continue
            day_number = int(item["day"])
            classes.append(
                {
                    "weekday": jalali_weekday_name(schedule_year, schedule_month, day_number),
                    "day": day_number,
                    "student": normalize_text(item.get("student")),
                }
            )
        classes.sort(key=lambda item: int(item["day"]))
        return classes

    def schedule_for_day(self, professor_name, jalali_day):
        return [item for item in self.schedule_for_professor(professor_name) if item["day"] == int(jalali_day)]


class ContactDirectory:
    def __init__(self, export_url):
        self.cache = WorkbookCache(export_url)

    def professor_phone_map(self):
        workbook = self.cache.workbook_data()
        worksheet = workbook[workbook.sheetnames[0]]
        result = {}
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            professor = normalize_text(row[1] if len(row) > 1 else "")
            phone = normalize_phone(row[2] if len(row) > 2 else "")
            if professor and phone:
                result[phone] = professor
        return result


def chat_registry():
    raw = read_json(CHAT_IDS_PATH, {})
    if isinstance(raw, list):
        return {str(chat_id): {"phone": None, "professor": None} for chat_id in raw}
    return {
        str(chat_id): {
            "phone": normalize_phone(value.get("phone")),
            "professor": normalize_text(value.get("professor")),
        }
        for chat_id, value in raw.items()
        if isinstance(value, dict)
    }


def save_chat_registry(registry):
    write_json(CHAT_IDS_PATH, registry)


def set_chat_professor(chat_id, phone, professor):
    registry = chat_registry()
    registry[str(chat_id)] = {"phone": normalize_phone(phone), "professor": normalize_text(professor)}
    save_chat_registry(registry)


def get_chat_professor(chat_id):
    return normalize_text(chat_registry().get(str(chat_id), {}).get("professor"))


def get_chat_registration(chat_id):
    item = chat_registry().get(str(chat_id), {})
    return {
        "phone": normalize_phone(item.get("phone")),
        "professor": normalize_text(item.get("professor")),
    }


def classes_keyboard():
    return {
        "keyboard": [[{"text": VIEW_CLASSES_TEXT}]],
        "resize_keyboard": True,
        "one_time_keyboard": False,
    }


def attendance_keyboard():
    return {
        "keyboard": [[{"text": ATTENDANCE_YES_TEXT}, {"text": ATTENDANCE_NO_TEXT}]],
        "resize_keyboard": True,
        "one_time_keyboard": True,
    }


def remove_keyboard():
    return {"remove_keyboard": True}


def format_phone(phone):
    return to_persian_digits(normalize_phone(phone))


def format_schedule(professor_name, schedule, schedule_reader):
    year, _, month_name = schedule_reader.month_info()
    if not schedule:
        return f"برای {professor_name} در این ماه برنامه‌ای ثبت نشده است."
    lines = [f"برنامه کلاس های {professor_name}:", ""]
    for index, item in enumerate(schedule, start=1):
        lines.append(
            f"{index}. {item['weekday']} {to_persian_digits(item['day'])} {month_name} {to_persian_digits(year)}\n"
            f"   دانشجو: {item['student']}"
        )
        lines.append("")
    lines.extend(["", SCHEDULE_FOOTER])
    return "\n".join(lines)


def format_login_schedule(professor_name, schedule, schedule_reader):
    return "\n\n".join(
        [
            LOGIN_SUCCESS_MESSAGE.format(professor_name=professor_name),
            format_schedule(professor_name, schedule, schedule_reader),
        ]
    )


def format_reminder(professor_name, day_schedule, schedule_reader):
    year, _, month_name = schedule_reader.month_info()
    day = day_schedule[0]["day"]
    lines = [
        f"یادآوری برنامه امروز {professor_name}",
        f"{day_schedule[0]['weekday']} {to_persian_digits(day)} {month_name} {to_persian_digits(year)}",
        "",
    ]
    for index, item in enumerate(day_schedule, start=1):
        lines.append(f"{index}. دانشجو: {item['student']}")
    return "\n".join(lines)


def format_attendance_question(item):
    return f"آیا دانشجوی {item['student']} در کلاس امروز شما حضور پیدا کرد؟"


def pending_attendance():
    return read_json(PENDING_ATTENDANCE_PATH, {})


def save_pending_attendance(value):
    write_json(PENDING_ATTENDANCE_PATH, value)


def local_absences():
    return read_json(ABSENCES_FALLBACK_PATH, [])


def append_local_absence(absence):
    items = local_absences()
    items.append(absence)
    write_json(ABSENCES_FALLBACK_PATH, items)


def append_absence_to_excel(absence):
    ABSENCES_XLSX_PATH.parent.mkdir(parents=True, exist_ok=True)
    if ABSENCES_XLSX_PATH.exists():
        workbook = load_workbook(ABSENCES_XLSX_PATH)
        worksheet = workbook.active
    else:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "غیبت دانشجویان"
        worksheet.append(["ردیف", "نام دانشجو", "نام استاد", "تاریخ غیبت"])

    next_number = max(1, worksheet.max_row)
    worksheet.append(
        [
            next_number,
            absence["student"],
            absence["professor"],
            absence["date"],
        ]
    )
    workbook.save(ABSENCES_XLSX_PATH)


def append_absence_to_webhook(config, absence):
    webhook_url = str(config.get("absence_webhook_url", "")).strip()
    if not webhook_url:
        return False
    response = requests.post(
        webhook_url,
        json={
            "student": absence["student"],
            "professor": absence["professor"],
            "date": absence["date"],
        },
        timeout=30,
    )
    response.raise_for_status()
    return True


def record_absence(config, absence):
    try:
        if append_absence_to_webhook(config, absence):
            return
    except Exception as exc:
        print(f"Apps Script absence write failed: {exc}")

    try:
        append_absence_to_excel(absence)
        return
    except Exception as exc:
        print(f"Excel absence write failed: {exc}")
    append_local_absence(absence)


def ask_next_attendance_question(bot, chat_id):
    pending = pending_attendance()
    item = pending.get(str(chat_id), {}).get("active")
    if not item:
        return False
    bot.send_message(chat_id, format_attendance_question(item), keyboard=attendance_keyboard())
    return True


def complete_attendance_answer(bot, config, chat_id, answer):
    pending = pending_attendance()
    state = pending.get(str(chat_id))
    if not state or not state.get("active"):
        return False

    active = state["active"]
    if answer == ATTENDANCE_NO_TEXT:
        record_absence(config, active)
        bot.send_message(chat_id, "غیبت دانشجو ثبت شد", keyboard=classes_keyboard())
    else:
        bot.send_message(chat_id, "با تشکر از شما", keyboard=classes_keyboard())

    queue = state.get("queue", [])
    if queue:
        state["active"] = queue.pop(0)
        state["queue"] = queue
        pending[str(chat_id)] = state
        save_pending_attendance(pending)
        ask_next_attendance_question(bot, chat_id)
    else:
        pending.pop(str(chat_id), None)
        save_pending_attendance(pending)
    return True


def handle_message(bot, config, schedule_reader, contact_directory, message):
    chat_id = (message.get("chat") or {}).get("id")
    text = normalize_text(message.get("text"))
    if chat_id is None:
        return

    if text in {ATTENDANCE_YES_TEXT, ATTENDANCE_NO_TEXT}:
        if complete_attendance_answer(bot, config, chat_id, text):
            return

    registration = get_chat_registration(chat_id)
    registered_phone = registration["phone"]
    registered_professor = registration["professor"]

    if text == "/start":
        if registered_phone and registered_professor:
            schedule = schedule_reader.schedule_for_professor(registered_professor)
            bot.send_message(
                chat_id,
                format_login_schedule(registered_professor, schedule, schedule_reader),
                keyboard=classes_keyboard(),
            )
            return
        bot.send_message(chat_id, START_MESSAGE)
        return

    if text == VIEW_CLASSES_TEXT:
        if registered_phone and registered_professor:
            schedule = schedule_reader.schedule_for_professor(registered_professor)
            bot.send_message(
                chat_id,
                format_schedule(registered_professor, schedule, schedule_reader),
                keyboard=classes_keyboard(),
            )
            return
        bot.send_message(chat_id, START_MESSAGE)
        return

    phone = normalize_phone(text)
    if registered_phone and registered_professor:
        if phone and phone != registered_phone:
            bot.send_message(
                chat_id,
                REGISTERED_PHONE_MESSAGE.format(phone=format_phone(registered_phone)),
                keyboard=classes_keyboard(),
            )
            return
        schedule = schedule_reader.schedule_for_professor(registered_professor)
        bot.send_message(
            chat_id,
            format_schedule(registered_professor, schedule, schedule_reader),
            keyboard=classes_keyboard(),
        )
        return

    phones = contact_directory.professor_phone_map()
    if phone in phones:
        professor = phones[phone]
        set_chat_professor(chat_id, phone, professor)
        schedule = schedule_reader.schedule_for_professor(professor)
        bot.send_message(
            chat_id,
            format_login_schedule(professor, schedule, schedule_reader),
            keyboard=classes_keyboard(),
        )
        return

    bot.send_message(chat_id, INVALID_PHONE_MESSAGE)


def reminder_key(chat_id, year, month, day):
    return f"{chat_id}:{year:04d}-{month:02d}-{day:02d}"


def send_due_reminders(bot, schedule_reader, target_date=None, dry_run=False):
    year, month, day = parse_jalali_date(target_date)
    sheet_year, sheet_month, _ = schedule_reader.month_info()
    if year != sheet_year or month != sheet_month:
        return []
    messages = []
    for chat_id, value in chat_registry().items():
        professor = normalize_text(value.get("professor"))
        if not professor:
            continue
        day_schedule = schedule_reader.schedule_for_day(professor, day)
        if not day_schedule:
            continue
        key = reminder_key(chat_id, year, month, day)
        if key in REMINDERS_SENT_THIS_RUN:
            continue
        text = format_reminder(professor, day_schedule, schedule_reader)
        messages.append({"chat_id": chat_id, "professor": professor, "text": text})
        if not dry_run:
            bot.send_message(chat_id, text)
            REMINDERS_SENT_THIS_RUN.add(key)
    return messages


def attendance_key(chat_id, year, month, day, student):
    return f"{chat_id}:{year:04d}-{month:02d}-{day:02d}:{compact_text(student)}"


def jalali_date_text(year, month, day):
    month_name = next((name for name, number in PERSIAN_MONTHS.items() if number == month), str(month))
    month_text = to_persian_digits(f"{month:02d}")
    day_text = to_persian_digits(f"{day:02d}")
    return f"{to_persian_digits(year)}/{month_text}/{day_text} - {month_name}"


def send_due_attendance_questions(bot, schedule_reader, target_date=None, dry_run=False):
    year, month, day = parse_jalali_date(target_date)
    sheet_year, sheet_month, _ = schedule_reader.month_info()
    if year != sheet_year or month != sheet_month:
        return []

    pending = pending_attendance()
    messages = []

    for chat_id, value in chat_registry().items():
        if str(chat_id) in pending:
            continue
        professor = normalize_text(value.get("professor"))
        if not professor:
            continue
        day_schedule = schedule_reader.schedule_for_day(professor, day)
        question_items = []
        for item in day_schedule:
            key = attendance_key(chat_id, year, month, day, item["student"])
            if key in ATTENDANCE_SENT_THIS_RUN:
                continue
            absence_item = {
                "professor": professor,
                "student": item["student"],
                "date": jalali_date_text(year, month, day),
                "key": key,
            }
            question_items.append(absence_item)
        if not question_items:
            continue
        messages.extend({"chat_id": chat_id, **item} for item in question_items)
        if not dry_run:
            pending[str(chat_id)] = {"active": question_items[0], "queue": question_items[1:]}
            bot.send_message(chat_id, format_attendance_question(question_items[0]), keyboard=attendance_keyboard())
            for item in question_items:
                ATTENDANCE_SENT_THIS_RUN.add(item["key"])

    if messages and not dry_run:
        save_pending_attendance(pending)
    return messages


def reminder_loop(bot, config, schedule_reader):
    reminder_time = config.get("reminder_time", "09:00")
    attendance_time = config.get("attendance_time", "21:00")
    target_date = config.get("current_jalali_date")
    interval_seconds = int(config.get("scheduler_interval_seconds", SCHEDULER_INTERVAL_SECONDS))
    timezone = get_bot_timezone(config.get("bot_timezone", BOT_TIMEZONE))
    print(
        "Scheduler is using "
        f"{timezone_label(timezone)}; now={datetime.now(timezone).strftime('%Y-%m-%d %H:%M:%S')}; "
        f"reminder={reminder_time}; attendance={attendance_time}; interval={interval_seconds}s"
    )
    while True:
        now = datetime.now(timezone)
        if time_matches(now, reminder_time):
            send_due_reminders(bot, schedule_reader, target_date=target_date)
        if time_matches(now, attendance_time):
            send_due_attendance_questions(bot, schedule_reader, target_date=target_date)
        time.sleep(interval_seconds)


def run():
    config = load_config()
    start_health_server()
    bot = BaleBot(config)
    schedule_reader = SheetSchedule(config["sheet_export_url"], config)
    contact_directory = ContactDirectory(config["contacts_export_url"])
    Thread(target=reminder_loop, args=(bot, config, schedule_reader), daemon=True).start()
    print("Ostad Yar bot is running. Press Ctrl+C to stop.")
    while True:
        try:
            response = bot.get_updates()
            for update in response.get("result", []):
                update_id = update.get("update_id")
                if update_id is not None:
                    bot.offset = int(update_id) + 1
                message = update.get("message")
                if message:
                    handle_message(bot, config, schedule_reader, contact_directory, message)
        except KeyboardInterrupt:
            print("\nBot stopped.")
            break
        except (urllib.error.URLError, urllib.error.HTTPError, requests.RequestException, TimeoutError) as exc:
            print(f"Network/API error: {exc}")
            time.sleep(5)
        except Exception as exc:
            print(f"Unexpected error: {exc}")
            time.sleep(5)


def preview_schedule(args):
    config = load_config()
    reader = SheetSchedule(config["sheet_export_url"], config)
    professor = ContactDirectory(config["contacts_export_url"]).professor_phone_map().get(normalize_phone(args.phone))
    if not professor:
        print(INVALID_PHONE_MESSAGE)
        return
    print(format_schedule(professor, reader.schedule_for_professor(professor), reader))


def preview_reminders(args):
    config = load_config()
    reader = SheetSchedule(config["sheet_export_url"], config)
    if args.phone:
        professor = ContactDirectory(config["contacts_export_url"]).professor_phone_map().get(normalize_phone(args.phone))
        if not professor:
            print(INVALID_PHONE_MESSAGE)
            return
        _, _, day = parse_jalali_date(args.date)
        day_schedule = reader.schedule_for_day(professor, day)
        if not day_schedule:
            print("No reminders found for this professor and date.")
            return
        print(format_reminder(professor, day_schedule, reader))
        return

    messages = send_due_reminders(None, reader, target_date=args.date, dry_run=True)
    if not messages:
        print("No reminders found for this date.")
        return
    for message in messages:
        print(f"CHAT_ID: {message['chat_id']}")
        print(message["text"])
        print("-" * 30)


def preview_attendance(args):
    config = load_config()
    reader = SheetSchedule(config["sheet_export_url"], config)
    messages = send_due_attendance_questions(None, reader, target_date=args.date, dry_run=True)
    if not messages:
        print("No attendance questions found for this date.")
        return
    for message in messages:
        print(f"CHAT_ID: {message['chat_id']}")
        print(format_attendance_question(message))
        print("-" * 30)


def main():
    parser = argparse.ArgumentParser()
    subparsers = parser.add_subparsers(dest="command")
    schedule_parser = subparsers.add_parser("preview-schedule")
    schedule_parser.add_argument("--phone", required=True)
    reminder_parser = subparsers.add_parser("test-reminders")
    reminder_parser.add_argument("--date", required=True, help="Jalali date, for example 1405-04-01")
    reminder_parser.add_argument("--phone", help="Optional professor phone for testing before Bale login")
    attendance_parser = subparsers.add_parser("test-attendance")
    attendance_parser.add_argument("--date", required=True, help="Jalali date, for example 1405-04-01")
    args = parser.parse_args()
    if args.command == "preview-schedule":
        preview_schedule(args)
    elif args.command == "test-reminders":
        preview_reminders(args)
    elif args.command == "test-attendance":
        preview_attendance(args)
    else:
        run()


if __name__ == "__main__":
    main()
