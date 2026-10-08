import os
import re
import time
from datetime import date
from io import BytesIO

import requests
from openpyxl import load_workbook


DEFAULT_STUDENT_CONTACTS_EXPORT_URL = (
    "https://docs.google.com/spreadsheets/d/"
    "1NLGwdUXK6cDLPx56gKe9oX89UbMSUhFwHHFA0uVDTt4/export?format=xlsx"
)
DEFAULT_STUDENT_CHAT_REGISTRY_EXPORT_URL = (
    "https://docs.google.com/spreadsheets/d/"
    "1pleT6E8_upnr_0TJImZXeENZD-mzdqcdqgQTUvCIVN4/export?format=xlsx"
)
DEFAULT_STUDENT_CHAT_REGISTRY_WEBHOOK_URL = (
    "https://script.google.com/macros/s/"
    "AKfycbzakE5t-mn_OlbeUJw9udcak8falohVf6XAno9tAOFNTpKWplxhFhdN5rewy80ekvg0UA/exec"
)

STUDENT_PHONE_REQUEST_MESSAGE = (
    "لطفا برای احراز هویت و مشاهده برنامه کلینیک ویژه، از دکمه اشتراک گذاری شماره همراه استفاده نمایید."
)
STUDENT_VIEW_CLASSES_TEXT = "دریافت برنامه ماه کلینیک ویژه من"
STUDENT_FLASHCARDS_TEXT = "دریافت فلش کارت ها"
STUDENT_EVALUATE_PROFESSORS_TEXT = "ارزیابی اساتید"
STUDENT_LOGBOOK_TEXT = "دریافت لاگ بوک اساتید"
STUDENT_ABSENCE_STATUS_TEXT = "مشاهده ی حضور و غیاب من"

STUDENT_REMINDERS_SENT_THIS_RUN = set()
STUDENT_SHEET_ROW_CACHE = {}
PERSIAN_MONTHS_BY_NUMBER = {
    1: "فروردین",
    2: "اردیبهشت",
    3: "خرداد",
    4: "تیر",
    5: "مرداد",
    6: "شهریور",
    7: "مهر",
    8: "آبان",
    9: "آذر",
    10: "دی",
    11: "بهمن",
    12: "اسفند",
}
PERSIAN_WEEKDAYS = [
    "دوشنبه",
    "سه شنبه",
    "چهارشنبه",
    "پنج شنبه",
    "جمعه",
    "شنبه",
    "یکشنبه",
]


def normalize_text(value):
    text = str(value or "").strip()
    text = text.replace("ي", "ی").replace("ك", "ک")
    return re.sub(r"\s+", " ", text)


def normalize_digits(value):
    return str(value or "").translate(str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789"))


def normalize_phone(value):
    if isinstance(value, float) and value.is_integer():
        value = str(int(value))
    elif isinstance(value, int):
        value = str(value)
    else:
        value = str(value or "").strip()
        if re.fullmatch(r"\d+\.0", value):
            value = value[:-2]
    value = normalize_digits(value)
    digits = re.sub(r"\D+", "", value)
    if digits.startswith("0098"):
        digits = "0" + digits[4:]
    elif digits.startswith("98") and len(digits) >= 12:
        digits = "0" + digits[2:12]
    elif digits.startswith("9") and len(digits) == 10:
        digits = "0" + digits
    return digits


def normalize_id(value):
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    text = normalize_text(value)
    if re.fullmatch(r"\d+\.0", text):
        return text[:-2]
    return text


def compact_text(value):
    return normalize_text(value).replace(" ", "")


def match_text_key(value):
    return compact_text(value).replace("دانسجو", "دانشجو")


def to_persian_digits(value):
    return str(value).translate(str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹"))


def jalali_to_gregorian(j_year, j_month, j_day):
    j_days_in_month = [31, 31, 31, 31, 31, 31, 30, 30, 30, 30, 30, 29]
    g_days_in_month = [31, 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]

    jy = j_year - 979
    jm = j_month - 1
    jd = j_day - 1

    j_day_no = 365 * jy + (jy // 33) * 8 + ((jy % 33) + 3) // 4
    for index in range(jm):
        j_day_no += j_days_in_month[index]
    j_day_no += jd

    g_day_no = j_day_no + 79
    gy = 1600 + 400 * (g_day_no // 146097)
    g_day_no %= 146097

    leap = True
    if g_day_no >= 36525:
        g_day_no -= 1
        gy += 100 * (g_day_no // 36524)
        g_day_no %= 36524
        if g_day_no >= 365:
            g_day_no += 1
        else:
            leap = False

    gy += 4 * (g_day_no // 1461)
    g_day_no %= 1461

    if g_day_no >= 366:
        leap = False
        g_day_no -= 1
        gy += g_day_no // 365
        g_day_no %= 365

    gm = 0
    while gm < 11:
        days_in_month = g_days_in_month[gm]
        if gm == 1 and leap:
            days_in_month += 1
        if g_day_no < days_in_month:
            break
        g_day_no -= days_in_month
        gm += 1

    return gy, gm + 1, g_day_no + 1


def jalali_weekday_name(j_year, j_month, j_day):
    g_year, g_month, g_day = jalali_to_gregorian(j_year, j_month, j_day)
    return PERSIAN_WEEKDAYS[date(g_year, g_month, g_day).weekday()]


def format_absence_date(date_text):
    date_text = normalize_digits(normalize_text(date_text))
    match = re.search(r"(\d{4})[/-](\d{1,2})[/-](\d{1,2})", date_text)
    if not match:
        return date_text
    year, month, day = (int(part) for part in match.groups())
    month_name = PERSIAN_MONTHS_BY_NUMBER.get(month)
    if not month_name:
        return date_text
    weekday = jalali_weekday_name(year, month, day)
    return f"{weekday} {to_persian_digits(day)} {month_name} {to_persian_digits(year)}"


def config_or_env(config, env_name, key, default):
    return os.getenv(env_name, config.get(key, default)).strip()


def cached_sheet_rows(export_url, cache_seconds=10):
    now = time.time()
    cached = STUDENT_SHEET_ROW_CACHE.get(export_url)
    if cached and now - cached["loaded_at"] <= cache_seconds:
        return cached["rows"]

    response = requests.get(export_url, timeout=30)
    response.raise_for_status()
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    headers = [normalize_text(cell.value) for cell in worksheet[1]]
    rows = []
    for raw_row in worksheet.iter_rows(min_row=2, values_only=True):
        row = {}
        for index, header in enumerate(headers):
            if header:
                row[header] = raw_row[index] if index < len(raw_row) else None
        if any(value not in (None, "") for value in row.values()):
            rows.append(row)
    STUDENT_SHEET_ROW_CACHE[export_url] = {"loaded_at": now, "rows": rows}
    return rows


def invalidate_sheet_cache(export_url):
    STUDENT_SHEET_ROW_CACHE.pop(export_url, None)


def post_sheet_action(webhook_url, payload):
    response = requests.post(webhook_url, json=payload, timeout=30)
    response.raise_for_status()
    try:
        result = response.json()
    except ValueError as exc:
        preview = response.text[:500] if response.text else "<empty response>"
        raise RuntimeError(f"Google Sheet webhook returned non-JSON response: {preview}") from exc
    if not result.get("ok"):
        raise RuntimeError(result.get("error") or "Google Sheet webhook returned ok=false")
    return result


class WorkbookCache:
    def __init__(self, export_url):
        self.export_url = export_url
        self.loaded_at = 0
        self.cache_seconds = int(os.getenv("SHEET_CACHE_SECONDS", "300"))
        self.workbook = None

    def workbook_data(self):
        now = time.time()
        if self.workbook is None or now - self.loaded_at > self.cache_seconds:
            response = requests.get(self.export_url, timeout=30)
            response.raise_for_status()
            self.workbook = load_workbook(BytesIO(response.content), data_only=True)
            self.loaded_at = now
        return self.workbook


class StudentDirectory:
    def __init__(self, export_url):
        self.cache = WorkbookCache(export_url)

    def header_index(self, headers, aliases, default):
        for alias in aliases:
            alias_key = compact_text(alias)
            if alias_key in headers:
                return headers.index(alias_key)
        return default

    def student_phone_map(self):
        workbook = self.cache.workbook_data()
        worksheet = workbook[workbook.sheetnames[0]]
        headers = [compact_text(cell.value) for cell in worksheet[1]]
        name_index = self.header_index(
            headers,
            [
                "student name",
                "نام و نام خانوادگی",
                "نام و نام خانوادگی خود را انتخاب کنید",
            ],
            1,
        )
        phone_index = self.header_index(
            headers,
            [
                "phone number",
                "شماره تلفن",
                "شماره همراه",
                "شماره تلفنی که عضو بله هستین به صورت اعداد انگلیسی",
            ],
            2,
        )
        result = {}
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            student = normalize_text(row[name_index] if len(row) > name_index else "")
            phone = normalize_phone(row[phone_index] if len(row) > phone_index else "")
            if student and phone:
                result[phone] = student
        return result


def normalize_student_registry_rows(rows):
    registry = {}
    for row in rows:
        chat_id = normalize_id(row.get("chat_id"))
        if not chat_id:
            continue
        registry[str(chat_id)] = {
            "phone": normalize_phone(row.get("phone")),
            "student": normalize_text(row.get("student")),
        }
    return registry


def student_chat_registry(config):
    export_url = config_or_env(
        config,
        "STUDENT_CHAT_REGISTRY_EXPORT_URL",
        "student_chat_registry_export_url",
        DEFAULT_STUDENT_CHAT_REGISTRY_EXPORT_URL,
    )
    return normalize_student_registry_rows(cached_sheet_rows(export_url))


def set_student_chat(chat_id, phone, student, config):
    export_url = config_or_env(
        config,
        "STUDENT_CHAT_REGISTRY_EXPORT_URL",
        "student_chat_registry_export_url",
        DEFAULT_STUDENT_CHAT_REGISTRY_EXPORT_URL,
    )
    webhook_url = config_or_env(
        config,
        "STUDENT_CHAT_REGISTRY_WEBHOOK_URL",
        "student_chat_registry_webhook_url",
        DEFAULT_STUDENT_CHAT_REGISTRY_WEBHOOK_URL,
    )
    post_sheet_action(
        webhook_url,
        {
            "action": "upsert_student_chat",
            "chat_id": normalize_id(chat_id),
            "phone": normalize_phone(phone),
            "student": normalize_text(student),
        },
    )
    invalidate_sheet_cache(export_url)


def get_student_registration(chat_id, config):
    item = student_chat_registry(config).get(normalize_id(chat_id), {})
    return {
        "phone": normalize_phone(item.get("phone")),
        "student": normalize_text(item.get("student")),
    }


def student_chat_id_by_name(student_name, config):
    student_key = match_text_key(student_name)
    for chat_id, value in student_chat_registry(config).items():
        if match_text_key(value.get("student")) == student_key:
            return chat_id
    return ""


def student_keyboard():
    return {
        "keyboard": [
            [{"text": STUDENT_VIEW_CLASSES_TEXT}],
            [{"text": STUDENT_FLASHCARDS_TEXT}, {"text": STUDENT_EVALUATE_PROFESSORS_TEXT}],
            [{"text": STUDENT_LOGBOOK_TEXT}, {"text": STUDENT_ABSENCE_STATUS_TEXT}],
        ],
        "resize_keyboard": True,
        "one_time_keyboard": False,
    }


def format_student_schedule(student_name, schedule, schedule_reader):
    year, _, month_name = schedule_reader.month_info()
    if not schedule:
        return f"برای {student_name} در ماه {month_name} برنامه‌ای ثبت نشده است."
    lines = [
        f"اینترن گرامی , {student_name}",
        f"برنامه ماه {month_name} کلینیک ویژه شما به صورت زیر است :",
        "",
    ]
    for index, item in enumerate(schedule, start=1):
        lines.append(
            f"{index}. {item['weekday']} {to_persian_digits(item['day'])} {month_name} {to_persian_digits(year)}\n"
            f"   استاد: {item['professor']}"
        )
        lines.append("")
    return "\n".join(lines).strip()


def format_student_reminder(student_name, day_schedule, schedule_reader):
    year, _, month_name = schedule_reader.month_info()
    day = day_schedule[0]["day"]
    lines = [
        f"با عرض سلام، یادآوری برنامه امروز {student_name}",
        f"{day_schedule[0]['weekday']} {to_persian_digits(day)} {month_name} {to_persian_digits(year)}",
        "",
    ]
    for index, item in enumerate(day_schedule, start=1):
        lines.append(f"{index}. امروز با استاد {item['professor']} کلاس دارید.")
    return "\n".join(lines)


def format_logbook_links():
    return "\n".join(
        [
            "‼️ لیست کامل لاگ‌بوک‌های کلینیک ویژه عصر",
            "",
            "🔹 اساتید اطفال",
            "https://t.me/Clinic_mui/108",
            "",
            "🔹 اساتید داخلی",
            "https://t.me/Clinic_mui/107",
            "",
            "🔹 اساتید جراحی",
            "https://t.me/Clinic_mui/118",
            "",
            "🔹 اساتید زنان",
            "https://t.me/Clinic_mui/109",
            "",
            "🔹 اساتید اورولوژی",
            "https://t.me/Clinic_mui/110",
            "",
            "🔹 اساتید نورولوژی",
            "https://t.me/Clinic_mui/111",
            "",
            "🔹 اساتید نوروسرجری",
            "https://t.me/Clinic_mui/112",
            "",
            "🔹 اساتید قلب",
            "https://t.me/Clinic_mui/113",
            "",
            "🔷 اساتید گوش، حلق و بینی",
            "https://t.me/Clinic_mui/114",
            "",
            "🔹 اساتید چشم پزشکی",
            "https://t.me/Clinic_mui/115",
            "",
            "🔹 اساتید روانپزشکی",
            "https://t.me/Clinic_mui/116",
            "",
            "🔹 اساتید ارتوپدی",
            "https://t.me/Clinic_mui/117",
            "",
            "🔹 اساتید عفونی",
            "https://t.me/Clinic_mui/119",
        ]
    )


def student_absence_history(config, student_name, absent_text):
    workbook = WorkbookCache(config["contacts_export_url"]).workbook_data()
    if "غیبت دانشجویان" in workbook.sheetnames:
        worksheet = workbook["غیبت دانشجویان"]
    else:
        worksheet = workbook[workbook.sheetnames[-1]]
    target_key = match_text_key(student_name)
    student_row = None
    for row in range(2, worksheet.max_row + 1):
        value = normalize_text(worksheet.cell(row, 3).value)
        if match_text_key(value) == target_key:
            student_row = row
            break
    if not student_row:
        return []

    items = []
    for col in range(4, worksheet.max_column + 1):
        date_text = normalize_digits(normalize_text(worksheet.cell(1, col).value))
        status_text = normalize_text(worksheet.cell(student_row, col).value)
        if not date_text or not status_text:
            continue
        parts = [normalize_text(part) for part in status_text.split("-", 1)]
        status = parts[0]
        professor = parts[1] if len(parts) > 1 else ""
        attendance_status = "غیبت" if status == absent_text else "حضور"
        items.append(
            {
                "date": format_absence_date(date_text),
                "professor": professor,
                "status": attendance_status,
            }
        )
    return items


def format_student_absence_history(config, student_name, absent_text, schedule_reader=None):
    items = student_absence_history(config, student_name, absent_text)
    if not items:
        return "برای شما حضور و غیابی ثبت نشده است."
    month_name = ""
    if schedule_reader is not None:
        _year, _month, month_name = schedule_reader.month_info()
    header_month = f" ماه {month_name}" if month_name else ""
    lines = [f"برنامه حضور و غیاب{header_month} شما به صورت زیر است :", ""]
    for index, item in enumerate(items, start=1):
        professor = item["professor"] or "نامشخص"
        lines.append(f"{index}- {item['date']}")
        lines.append(f"کلینیک {professor} - {item['status']}")
        lines.append("")
    return "\n".join(lines).strip()


def student_reminder_key(chat_id, year, month, day):
    return f"student:{chat_id}:{year:04d}-{month:02d}-{day:02d}"


def send_due_student_reminders(bot, config, schedule_reader, parse_jalali_date_func, target_date=None, dry_run=False):
    year, month, day = parse_jalali_date_func(target_date)
    sheet_year, sheet_month, _ = schedule_reader.month_info()
    if year != sheet_year or month != sheet_month:
        return []
    messages = []
    for chat_id, value in student_chat_registry(config).items():
        student = normalize_text(value.get("student"))
        if not student:
            continue
        day_schedule = schedule_reader.schedule_for_student_day(student, day)
        if not day_schedule:
            continue
        key = student_reminder_key(chat_id, year, month, day)
        if key in STUDENT_REMINDERS_SENT_THIS_RUN:
            continue
        text = format_student_reminder(student, day_schedule, schedule_reader)
        messages.append({"chat_id": chat_id, "student": student, "text": text})
        if not dry_run:
            bot.send_message(chat_id, text, keyboard=student_keyboard())
            STUDENT_REMINDERS_SENT_THIS_RUN.add(key)
    return messages
